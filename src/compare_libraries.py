"""Find Apple Music tracks that are missing from an offline music library.

Compares an Apple Music XML export with an XML export of the offline library (for example
the purchased and downloaded folders combined in one iTunes library) and writes an Excel
file listing the Apple Music tracks with no offline copy. Both XML files are read-only
inputs; the workbook is written inside the workspace output directory.

Tracks are matched on title, artist and duration, ignoring the album (see
track_matching.py), so a song that sits on a different album or compilation still counts.
"""

import argparse
import os
import plistlib
import sys
from pathlib import Path
from urllib.parse import unquote, urlparse
import re

from config import DEFAULTS, load, resolve_path
from normalization import prepare_apple, prepare_offline
from track_matching import LOOSE_SECONDS, LooseIndex
from track_matching import POSSIBLE as WEAK
from track_matching import prepare as prepare_query

# Store-link fields are not guaranteed to exist; any of these that hold a URL are used.
LINK_FIELDS = ("Store URL", "iTunes URL", "Apple Music URL", "URL", "Track URL", "Album URL")
COLUMNS = (
    "Title",
    "Artist",
    "Album",
    "Year",
    "Duration",
    "iTunes Link",
    "Status",
    "Closest Offline Track",
    "Offline File",
    "Reason",
)
MATCHED, POSSIBLE, NOT_FOUND = "IN OFFLINE LIBRARY", "POSSIBLE MATCH", "MISSING"


def read_tracks(path):
    with open(path, "rb") as f:
        document = plistlib.load(f)
    tracks = document.get("Tracks") if isinstance(document, dict) else None
    if not isinstance(tracks, dict):
        raise ValueError(f"{path} is not an iTunes/Apple Music library XML with a Tracks section")
    return list(tracks.values())


def location_to_path(url):
    """Turn an iTunes file URL into a native path (Windows drive and UNC forms included)."""
    parsed = urlparse(url)
    path = unquote(parsed.path)
    if re.match(r"^/[A-Za-z]:/", path):
        return path[1:]
    if parsed.netloc not in ("", "localhost"):
        return "//" + parsed.netloc + path
    return path


def real_files(tracks):
    """Keep tracks that are files. Returns (kept, cloud_only_count, vanished_count or None).

    Cloud-only entries have no Location. When this computer can see most of the listed
    files, entries whose file no longer exists are dropped too; otherwise (the library
    lives on another computer) the listing is trusted and vanished is None.
    """
    listed = [t for t in tracks if t.get("Location") and t.get("Track Type") != "Remote"]
    cloud = len(tracks) - len(listed)
    step = max(1, len(listed) // 200)
    sample = listed[::step]
    visible = bool(sample) and (
        sum(os.path.isfile(location_to_path(t["Location"])) for t in sample) * 2 >= len(sample)
    )
    if not visible:
        return listed, cloud, None
    kept = [t for t in listed if os.path.isfile(location_to_path(t["Location"]))]
    return kept, cloud, len(listed) - len(kept)


def as_file_rows(tracks):
    """Present XML tracks as file-like rows so the track matcher can search them."""
    rows = []
    for number, track in enumerate(tracks):
        location = track.get("Location")
        rows.append(
            {
                "Full Path": location_to_path(location)
                if location
                else f"xml-track:{track.get('Persistent ID') or number}",
                "Library Root": "xml",
                "Title": track.get("Name", ""),
                "Artist": track.get("Artist", ""),
                "Album": track.get("Album", ""),
                "Album Artist": track.get("Album Artist", ""),
                "Duration": (track.get("Total Time") or 0) / 1000 or None,
                "Track Number": track.get("Track Number"),
                "Disc Number": track.get("Disc Number"),
                "Year": track.get("Year"),
                "Archive Eligible": "protected" not in str(track.get("Kind", "")).casefold(),
                "_track": track,
                "_virtual": not location,
            }
        )
    return rows


def link_for(track):
    for field in LINK_FIELDS:
        value = track.get(field)
        if isinstance(value, str) and value.startswith(("http://", "https://", "itms")):
            return value
    return ""


def duration_text(track):
    ms = track.get("Total Time")
    return f"{int(ms // 60000)}:{int(ms // 1000 % 60):02d}" if ms else ""


def compare(
    apple_tracks, offline_tracks, config=None, loose_seconds=LOOSE_SECONDS, include_matched=False
):
    """Return row dictionaries for Apple Music tracks, sorted for browsing."""
    config = config or DEFAULTS
    rows = as_file_rows(offline_tracks)
    for row in rows:
        prepare_offline(row)
    finder = LooseIndex(rows)
    results = []
    for original in apple_tracks:
        track = dict(original)
        prepare_apple(track)
        prepare_query(track)
        found = finder.find(track, config, loose_seconds)
        best = found[0] if found else None
        status = NOT_FOUND if not best else POSSIBLE if best["tier"] == WEAK else MATCHED
        if status == MATCHED and not include_matched:
            continue
        b = best["file"] if best else None
        results.append(
            {
                "Title": track.get("Name", ""),
                "Artist": track.get("Artist", ""),
                "Album": track.get("Album", ""),
                "Year": track.get("Year", ""),
                "Duration": duration_text(track),
                "iTunes Link": link_for(track),
                "Status": status,
                "Closest Offline Track": (
                    f"{b.get('Title', '')} - {b.get('Artist', '')} ({b.get('Album', '')})"
                    if b
                    else ""
                ),
                "Offline File": "" if not b or b.get("_virtual") else b["Full Path"],
                "Reason": (
                    f"{best['tier'].title()} match: {best['evidence']}"
                    if best
                    else "no offline track with this title and artist"
                ),
                "_sort": (
                    str(track.get("Artist", "")).casefold(),
                    str(track.get("Album", "")).casefold(),
                    track.get("Disc Number", 0),
                    track.get("Track Number", 0),
                    str(track.get("Name", "")).casefold(),
                ),
            }
        )
    results.sort(key=lambda r: r["_sort"])
    return results


def write_xlsx(rows, destination, columns=COLUMNS):
    """Write the table as a workbook: every column is its own cell, text stays literal."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(list(columns))
    for row in rows:
        values = []
        for column in columns:
            value = row.get(column, "")
            if isinstance(value, str):
                value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
            values.append(value)
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # literal text, never a formula
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index_, column in enumerate(columns, 1):
        longest = max([len(str(column))] + [len(str(r.get(column, ""))) for r in rows[:500]])
        ws.column_dimensions[ws.cell(1, index_).column_letter].width = min(max(longest + 2, 10), 60)
    temporary = destination.with_name(destination.stem + ".tmp.xlsx")
    wb.save(temporary)
    os.replace(temporary, destination)


def candidate_xmls(workspace):
    found = []
    for folder in (Path.cwd(), workspace, workspace / "input"):
        if folder.is_dir():
            for path in sorted(folder.glob("*.xml")):
                if path.resolve() not in found:
                    found.append(path.resolve())
    return found


def choose_xml(label, workspace, ask=input, tell=print, exclude=None):
    """Pick from XML files found nearby, or type/drag in a path."""
    options = [p for p in candidate_xmls(workspace) if p != exclude]
    tell(f"\nChoose the {label} XML.")
    for number, path in enumerate(options, 1):
        tell(f"  {number}. {path}")
    while True:
        hint = "number or path" if options else "path"
        raw = ask(f"{label} XML ({hint}): ").strip().strip('"').strip("'")
        if raw.isdigit() and 1 <= int(raw) <= len(options):
            return options[int(raw) - 1]
        if raw:
            path = resolve_path(raw, workspace)
            if path.is_file():
                return path
        tell("That is not an existing file. Try again.")


def ask_folders(ask=input, tell=print):
    tell("\nEnter the downloaded-music folder, one per answer. Enter done when finished.")
    folders = []
    while True:
        raw = ask(f"Music folder {len(folders) + 1}, or done: ").strip().strip('"').strip("'")
        if raw.casefold() in ("done", "") and folders:
            return folders
        if raw and Path(raw).expanduser().is_dir():
            folders.append(str(Path(raw).expanduser()))
        else:
            tell("That folder does not exist. Try again." if raw else "Add at least one folder.")


def output_path(c, requested=None, default_name="Apple_Music_Missing_From_Offline.xlsx"):
    """Workbook destination: inside the workspace output folder, never over an input or music root."""
    workspace = Path(c["_workspace"])
    default = Path(c["output_excel"]).parent / default_name
    path = resolve_path(requested, workspace) if requested else default
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Output must end in .xlsx")
    if not path.is_relative_to(workspace):
        raise ValueError("The workbook must be written inside the workspace")
    if any(path.is_relative_to(Path(root)) for root in c["offline_music_roots"]):
        raise ValueError("The workbook cannot be written inside a read-only music folder")
    return path


def main(argv=None, ask=input, tell=print):
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(
        description="Compare an Apple Music library XML with an offline-library XML and "
        "write an Excel file of the Apple Music tracks that have no offline copy"
    )
    parser.add_argument("--apple", help="Apple Music library XML (asked for if omitted)")
    parser.add_argument(
        "--offline",
        "--purchased",
        dest="offline",
        help="Offline music library XML (asked for if omitted)",
    )
    parser.add_argument(
        "--output", help="Excel .xlsx path inside the workspace (default: output/*.xlsx)"
    )
    parser.add_argument("--workspace", help="Workspace directory (default: project folder)")
    parser.add_argument(
        "--config", help="Local JSON settings path (default: workspace/config.json)"
    )
    parser.add_argument(
        "--include-matched",
        action="store_true",
        help="Also list tracks found in the offline library (Status = IN OFFLINE LIBRARY)",
    )
    parser.add_argument(
        "--duration-seconds",
        type=float,
        default=LOOSE_SECONDS,
        help="Largest duration difference still treated as the same song (default: %(default)s)",
    )
    args = parser.parse_args(argv)
    if not 0 <= args.duration_seconds <= 30:
        parser.error("--duration-seconds must be between 0 and 30")
    c = load(args.config, args.workspace)
    workspace = Path(c["_workspace"])
    interactive = args.apple is None or args.offline is None
    if interactive and not sys.stdin.isatty() and ask is input:
        parser.error("Supply --apple and --offline when not running in a terminal.")
    apple = (
        resolve_path(args.apple, workspace)
        if args.apple
        else choose_xml("Apple Music library", workspace, ask, tell)
    )
    offline = (
        resolve_path(args.offline, workspace)
        if args.offline
        else choose_xml("offline music library", workspace, ask, tell, exclude=apple)
    )
    for path in (apple, offline):
        if not path.is_file():
            parser.error(f"XML file not found: {path}")
    if apple == offline:
        parser.error("Choose two different XML files.")
    destination = output_path(c, args.output)
    if destination in (apple, offline):
        parser.error("The workbook cannot overwrite an input XML.")
    apple_tracks, offline_tracks = read_tracks(apple), read_tracks(offline)
    present, cloud_only, vanished = real_files(offline_tracks)
    rows = compare(apple_tracks, present, c, args.duration_seconds, args.include_matched)
    write_xlsx(rows, destination)
    counts = {s: sum(r["Status"] == s for r in rows) for s in (MATCHED, POSSIBLE, NOT_FOUND)}
    tell(f"\nApple Music tracks:      {len(apple_tracks):,}")
    tell(f"Offline library tracks: {len(offline_tracks):,} listed, {len(present):,} used")
    if cloud_only:
        tell(f"  ignored {cloud_only:,} cloud-only entries that have no file")
    if vanished:
        tell(f"  ignored {vanished:,} entries whose file no longer exists on this computer")
    if vanished is None:
        tell("  (file existence was not checked: the library files are not on this computer)")
    if args.include_matched:
        tell(f"In offline library:     {counts[MATCHED]:,}")
    else:
        tell(f"In offline library:     {len(apple_tracks) - len(rows):,}  (not listed)")
    tell(f"Possible match:         {counts[POSSIBLE]:,}  (check these by hand)")
    tell(f"Missing:                {counts[NOT_FOUND]:,}")
    if not any(r["iTunes Link"] for r in rows):
        tell("No iTunes/store link fields exist in this XML, so that column is empty.")
    tell(f"Excel file written: {destination}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
