from pathlib import Path
import plistlib
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import reconcile_libraries as rl
from config import DEFAULTS

CONFIG = {**DEFAULTS, "_workspace": ".", "_config_path": "config.json"}


def apple(pid, name, artist, seconds=3, **extra):
    return {
        "Persistent ID": pid,
        "Name": name,
        "Artist": artist,
        "Album": "Album",
        "Total Time": int(seconds * 1000),
        **extra,
    }


def offline(path, title, artist, seconds=3, **extra):
    return {
        "Full Path": str(path),
        "Library Root": str(Path(path).parent),
        "Title": title,
        "Artist": artist,
        "Album": "Album",
        "Duration": seconds,
        "Extension": ".mp3",
        "Codec": "mp3",
        "Bitrate": "320000",
        "File Size": 7_000_000,
        "Archive Eligible": True,
        "Integrity": "PASS",
        "Metadata Issues": "",
        **extra,
    }


def state(table):
    return lambda track: table.get(track["Persistent ID"], ("NOT_DOWNLOADED", None))


class ReconcileTests(unittest.TestCase):
    def test_location_urls_become_native_paths(self):
        self.assertEqual(
            rl.location_to_path("file://localhost/D:/iTunes/Music/A%20B/01%20Song.m4a"),
            "D:/iTunes/Music/A B/01 Song.m4a",
        )
        self.assertEqual(rl.location_to_path("file:///home/me/x.m4a"), "/home/me/x.m4a")
        self.assertEqual(rl.location_to_path("file://nas/share/x.m4a"), "//nas/share/x.m4a")

    def test_only_unchanged_local_purchases_make_a_copy_removable(self):
        purchased = [
            apple("P1", "Local", "Band"),
            apple("P2", "Cloud", "Band"),
            apple("P3", "Moved", "Band"),
            apple("P4", "Edited", "Band"),
        ]
        files = [offline(f"/net/{t['Name']}.mp3", t["Name"], "Band") for t in purchased]
        files.append(offline("/net/unrelated.mp3", "Unrelated", "Someone"))
        table = {
            "P1": ("OK", "/buy/local.m4a"),
            "P3": ("MISSING_LOCAL", "/buy/moved.m4a"),
            "P4": ("CHANGED", "/buy/edited.m4a"),
        }
        plan, albums, _, _ = rl.reconcile([], purchased, files, CONFIG, state(table))
        got = {Path(r["Internet File"]).stem: r["Recommendation"] for r in plan}
        self.assertEqual(
            got,
            {
                "Local": rl.REMOVE,
                "Cloud": rl.KEEP_REMOTE,
                "Moved": rl.KEEP_GONE,
                "Edited": rl.REVIEW_CHANGED,
            },
        )
        self.assertEqual(albums[0]["Purchased Tracks"], 4)
        self.assertEqual(albums[0]["With Internet Copy"], 4)
        self.assertEqual(albums[0]["Internet Files To Remove"], 1)

    def test_uncertain_files_are_never_removable(self):
        purchased = [apple("P1", "Signal", "Band")]
        files = [
            offline("/net/untagged.mp3", "", ""),
            offline("/net/live.mp3", "Signal (Live)", "Band"),
            offline("/net/long.mp3", "Signal", "Band", seconds=30),
            offline("/net/protected.m4p", "Signal", "Band", **{"Archive Eligible": False}),
        ]
        plan, _, _, _ = rl.reconcile([], purchased, files, CONFIG, state({"P1": ("OK", "/b.m4a")}))
        self.assertNotIn(rl.REMOVE, {r["Recommendation"] for r in plan})

    def test_purchased_files_inside_the_scanned_folder_are_protected(self):
        with tempfile.TemporaryDirectory() as td:
            kept = Path(td) / "Signal.m4a"
            purchased = [apple("P1", "Signal", "Band", Location=kept.as_uri())]
            files = [offline(kept, "Signal", "Band")]
            plan, _, _, protected = rl.reconcile([], purchased, files, CONFIG)
            self.assertEqual((plan, protected), ([], 1))

    def test_missing_everywhere_checks_both_places(self):
        purchased = [apple("P1", "Bought", "Band")]
        files = [offline("/net/found.mp3", "Downloaded", "Band")]
        tracks = [
            apple("A1", "Bought", "Band"),
            apple("A2", "Downloaded", "Band"),
            apple("A3", "Nowhere", "Nobody"),
            apple("A4", "Bought", "Band", seconds=40),
        ]
        _, _, missing, _ = rl.reconcile(tracks, purchased, files, CONFIG)
        got = {r["Title"]: r["Status"] for r in missing}
        self.assertEqual(got, {"Nowhere": rl.MISSING, "Bought": rl.POSSIBLE})


@unittest.skipUnless(
    shutil.which("ffmpeg") and shutil.which("ffprobe"), "Requires FFmpeg and FFprobe"
)
class ReconcileCLITests(unittest.TestCase):
    def make(self, path, title, artist):
        path.parent.mkdir(parents=True, exist_ok=True)
        subprocess.run(
            ["ffmpeg", "-nostdin", "-v", "error", "-f", "lavfi", "-i", "sine=duration=3"]
            + ["-c:a", "aac", "-metadata", f"title={title}", "-metadata", f"artist={artist}"]
            + ["-metadata", "album=Album", str(path)],
            check=True,
        )
        return path

    def test_end_to_end_plan_is_read_only(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            bought = self.make(base / "bought" / "Signal.m4a", "Signal", "Band")
            net = base / "net"
            files = [
                self.make(net / "Signal copy.m4a", "Signal", "Band"),
                self.make(net / "Cloud.m4a", "Cloud Song", "Band"),
                self.make(net / "Elsewhere.m4a", "Elsewhere", "Band"),
            ]
            store = [
                apple("P1", "Signal", "Band", Location=bought.as_uri(), Size=bought.stat().st_size),
                apple("P2", "Cloud Song", "Band"),
            ]
            streaming = [
                apple("A1", "Signal", "Band"),
                apple("A2", "Elsewhere", "Band"),
                apple("A3", "Nowhere", "Nobody"),
            ]
            for name, tracks in (("apple.xml", streaming), ("purchased.xml", store)):
                (base / name).write_bytes(
                    plistlib.dumps({"Tracks": {str(i): t for i, t in enumerate(tracks)}})
                )
            before = {p: p.read_bytes() for p in [bought, *files]}
            work = base / "work"
            result = subprocess.run(
                [sys.executable, str(Path(rl.__file__))]
                + ["--workspace", str(work), "--apple", str(base / "apple.xml")]
                + ["--purchased", str(base / "purchased.xml"), "--music-root", str(net)],
                text=True,
                encoding="utf-8",
                capture_output=True,
                timeout=120,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(before, {p: p.read_bytes() for p in before})

            def rows(name):
                sheet = load_workbook(work / "output" / name).active
                header, *body = sheet.iter_rows(values_only=True)
                self.assertGreater(sheet.max_column, 5)
                return [dict(zip(header, row)) for row in body]

            plan = {
                Path(r["Internet File"]).name: r["Recommendation"]
                for r in rows("Internet_Duplicates_Of_Purchased.xlsx")
            }
            self.assertEqual(plan, {"Signal copy.m4a": rl.REMOVE, "Cloud.m4a": rl.KEEP_REMOTE})
            self.assertEqual(
                [r["Title"] for r in rows("Apple_Music_Missing_Everywhere.xlsx")], ["Nowhere"]
            )
            self.assertEqual(len(rows("Purchased_Albums_In_Internet_Collection.xlsx")), 1)


if __name__ == "__main__":
    unittest.main()
