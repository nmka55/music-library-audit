"""Validate the completed report against the XML and cached source file stats."""

from collections import Counter
import datetime
import json
from pathlib import Path
import plistlib
import sqlite3
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from config import load


def records(sheet):
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    return [dict(zip(headers, row)) for row in iterator]


def main():
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workspace")
    parser.add_argument("--config")
    args = parser.parse_args()
    c = load(args.config, args.workspace)
    wb = load_workbook(c["output_excel"], read_only=True, data_only=False)
    expected = [
        "Dashboard",
        "Apple Music Library",
        "Offline Library",
        "Match Results",
        "Album Audit",
        "Duplicates",
    ]
    assert wb.sheetnames == expected, wb.sheetnames
    with open(c["apple_music_xml"], "rb") as f:
        original = list(plistlib.load(f)["Tracks"].values())
    apple = records(wb["Apple Music Library"])
    assert len(apple) == len(original)
    for source, export in zip(original, apple):
        for key, value in source.items():
            if isinstance(value, (str, int, float, bool, datetime.datetime)):
                assert export[key] == value, (source.get("Persistent ID"), key, value, export[key])
    offline = records(wb["Offline Library"])
    matches = records(wb["Match Results"])
    assert len(matches) == len(original)
    assert len({r["Full Path"] for r in offline}) == len(offline)
    by_path = {r["Full Path"]: r for r in offline}
    counts = Counter(r["Status"] for r in matches)
    assert set(counts) <= {"MATCHED", "MISSING", "REVIEW"}
    for row in matches:
        if row["Status"] == "MATCHED":
            assert row["Archive Eligible"] is True, row
            assert row["Matched File Path"] in by_path, row
            assert by_path[row["Matched File Path"]]["Archive Eligible"] is True, row
            if row["Match Type"] == "AUTOMATIC":
                assert row["Overall Confidence"] >= c["matched_threshold"], row
                assert row["Version Conflict"] is False, row
                assert row["Duration Difference"] <= c["duration_tolerance_seconds"], row
                assert row["Title Score"] >= 97 and row["Artist Score"] >= 97, row
    summary = {r["Metric"]: r["Value"] for r in records(wb["Dashboard"])}
    assert summary["Total Apple Music tracks"] == len(original)
    assert summary["Supported-extension files discovered"] == len(offline)
    assert summary["Total offline audio files"] == sum(
        r.get("File Type") != "AppleDouble metadata sidecar" for r in offline
    )
    for status in ("MATCHED", "REVIEW", "MISSING"):
        assert summary[status] == counts[status]
    row_counts = {ws.title: ws.max_row - 1 for ws in wb}
    wb.close()
    # Check presentation metadata without loading a million styled cells in memory.
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(c["output_excel"]) as archive:
        for i in range(1, 7):
            tree = ET.fromstring(archive.read(f"xl/worksheets/sheet{i}.xml"))
            assert tree.find("x:autoFilter", ns) is not None
            assert tree.find("x:sheetViews/x:sheetView/x:pane", ns).attrib["topLeftCell"] == "A2"
            assert tree.find(".//x:f", ns) is None, "Unexpected formula in imported metadata"
            if i == 3 and offline or i == 4 and any(r.get("Matched File Path") for r in matches):
                assert tree.find("x:hyperlinks", ns) is not None
    changed = []
    with sqlite3.connect(Path(c["cache_directory"]) / "inventory.sqlite3") as db:
        stats = {
            path: (size, mtime)
            for path, size, mtime in db.execute("SELECT path,size,mtime FROM files")
        }
    for path in by_path:
        try:
            st = Path(path).stat()
            if stats.get(path) != (st.st_size, st.st_mtime_ns):
                changed.append(path)
        except OSError:
            changed.append(path)
    result = {
        "validated_at": datetime.datetime.now().isoformat(),
        "workbook": c["output_excel"],
        "sheet_rows": row_counts,
        "match_counts": dict(counts),
        "source_files_changed_or_missing_since_scan": changed,
        "status": "PASS" if not changed else "SOURCE CHANGED",
    }
    Path(c["output_excel"]).with_name("validation.json").write_text(
        json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps(result, indent=2, ensure_ascii=False))
    assert not changed, "Some source files changed or disappeared after inventory"


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
