from pathlib import Path
import plistlib
import sys
import tempfile
import unittest

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import compare_libraries as cl


def track(pid, name, artist, seconds, **extra):
    return {
        "Persistent ID": pid,
        "Name": name,
        "Artist": artist,
        "Album": "Album",
        "Year": 2001,
        "Total Time": int(seconds * 1000),
        **extra,
    }


def write_xml(path, tracks):
    path.write_bytes(plistlib.dumps({"Tracks": {str(i): t for i, t in enumerate(tracks)}}))


class CompareTests(unittest.TestCase):
    def statuses(self, apple, purchased, **kwargs):
        rows = cl.compare(apple, purchased, include_matched=True, **kwargs)
        return {r["Title"]: r["Status"] for r in rows}

    def test_gates_are_conservative(self):
        purchased = [
            track("B1", "Signal", "Band", 200),
            track("B2", "Echo", "Band", 200),
            track("B3", "Orbit (Live)", "Band", 200),
        ]
        apple = [
            track("A1", "Signal", "Band", 202),  # within tolerance
            track("A2", "Echo", "Band", 230),  # duration differs
            track("A3", "Orbit", "Band", 200),  # version differs
            track("A4", "Unknown Song", "Band", 200),  # absent
        ]
        result = self.statuses(apple, purchased)
        self.assertEqual(result["Signal"], cl.MATCHED)
        self.assertEqual(result["Echo"], cl.POSSIBLE)
        self.assertEqual(result["Orbit"], cl.POSSIBLE)
        self.assertEqual(result["Unknown Song"], cl.NOT_FOUND)

    def test_track_id_alone_never_matches_and_missing_duration_is_uncertain(self):
        purchased = [dict(track("B1", "Different", "Other", 100), **{"Track ID": 7})]
        apple = [dict(track("A1", "Signal", "Band", 100), **{"Track ID": 7})]
        self.assertEqual(self.statuses(apple, purchased), {"Signal": cl.NOT_FOUND})
        purchased = [track("B1", "Signal", "Band", 100)]
        no_time = [{"Persistent ID": "A", "Name": "Signal", "Artist": "Band"}]
        self.assertEqual(self.statuses(no_time, purchased), {"Signal": cl.POSSIBLE})

    def test_featured_credit_formatting_and_default_hides_matches(self):
        purchased = [track("B1", "Song (feat. Guest)", "Band", 100)]
        apple = [track("A1", "Song", "Band & Guest", 100)]
        self.assertEqual(cl.compare(apple, purchased), [])
        self.assertEqual(self.statuses(apple, purchased), {"Song": cl.MATCHED})

    def test_links_are_used_only_when_present(self):
        apple = [
            track("A1", "One", "Band", 10, **{"Store URL": "https://example.test/one"}),
            track("A2", "Two", "Band", 10),
        ]
        rows = {r["Title"]: r for r in cl.compare(apple, [])}
        self.assertEqual(rows["One"]["iTunes Link"], "https://example.test/one")
        self.assertEqual(rows["Two"]["iTunes Link"], "")

    def test_cli_writes_multicolumn_workbook_without_changing_inputs(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            apple, offline = base / "apple.xml", base / "offline.xml"
            music = base / "Signal.m4a"
            music.write_bytes(b"audio")
            write_xml(
                apple,
                [
                    track("A1", "Signal", "Band, The", 200),
                    track("A2", "=Boom", "X", 60),
                    track("A3", "Cloud Song", "Band", 100),
                ],
            )
            write_xml(
                offline,
                [
                    track("B1", "Signal", "Band, The", 200, Location=music.as_uri()),
                    track("B2", "Cloud Song", "Band", 100, **{"Track Type": "Remote"}),
                ],
            )
            before = (apple.read_bytes(), offline.read_bytes())
            messages = []
            code = cl.main(
                ["--workspace", str(base), "--apple", str(apple), "--offline", str(offline)],
                tell=messages.append,
            )
            self.assertEqual(code, 0)
            self.assertEqual(before, (apple.read_bytes(), offline.read_bytes()))
            self.assertTrue(any("cloud-only" in m for m in messages))
            sheet = load_workbook(base / "output" / "Apple_Music_Missing_From_Offline.xlsx").active
            table = list(sheet.iter_rows(values_only=True))
            # Every field sits in its own cell, never one cell holding a whole row.
            self.assertEqual(sheet.max_column, len(cl.COLUMNS))
            self.assertEqual(list(table[0]), list(cl.COLUMNS))
            rows = {r[0]: dict(zip(table[0], r)) for r in table[1:]}
            self.assertEqual(set(rows), {"=Boom", "Cloud Song"})
            self.assertEqual(sheet["A2"].data_type, "s")
            self.assertEqual(rows["=Boom"]["Duration"], "1:00")
            self.assertEqual(rows["=Boom"]["Year"], 2001)
            self.assertEqual(rows["Cloud Song"]["Status"], cl.NOT_FOUND)

    def test_real_files_drops_cloud_only_and_vanished_entries_when_files_are_visible(self):
        with tempfile.TemporaryDirectory() as td:
            kept = Path(td) / "kept.m4a"
            kept.write_bytes(b"x")
            gone = Path(td) / "gone.m4a"
            tracks = [
                track("1", "Kept", "A", 1, Location=kept.as_uri()),
                track("2", "Gone", "A", 1, Location=gone.as_uri()),
                track("3", "Cloud", "A", 1),
            ]
            present, cloud, vanished = cl.real_files(tracks[:1] + tracks[2:])
            self.assertEqual(([t["Name"] for t in present], cloud, vanished), (["Kept"], 1, 0))
            present, cloud, vanished = cl.real_files([tracks[1]])
            self.assertEqual((len(present), vanished), (1, None))  # nothing visible: trust XML

    def test_interactive_choice_by_number_or_path_and_output_boundary(self):
        with tempfile.TemporaryDirectory() as td, tempfile.TemporaryDirectory() as outside:
            base = Path(td)
            first, second = base / "a.xml", base / "b.xml"
            write_xml(first, [track("A1", "One", "Band", 10)])
            write_xml(second, [track("B1", "Two", "Band", 10)])
            answers = iter(["nope.xml", str(first), "1"])
            tell = lambda _: None
            chosen = cl.choose_xml("Apple Music library", base, lambda _: next(answers), tell)
            self.assertEqual(chosen, first.resolve())
            other = cl.choose_xml(
                "offline", base, lambda _: next(answers), tell, exclude=first.resolve()
            )
            self.assertEqual(other, second.resolve())
            with self.assertRaises(ValueError):
                cl.output_path(cl.load(workspace=base), str(Path(outside) / "x.xlsx"))
            with self.assertRaises(ValueError):
                cl.output_path(cl.load(workspace=base), "output/x.csv")


if __name__ == "__main__":
    unittest.main()
