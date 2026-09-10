import datetime
import json
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

COLORS = {"MATCHED": "E2F0D9", "REVIEW": "FCE4D6", "MISSING": "F4CCCC", "WARNING": "FFF2CC"}


def cell_value(value):
    if isinstance(value, (list, dict, set, tuple)):
        value = json.dumps(
            list(value) if isinstance(value, set) else value, ensure_ascii=False, default=str
        )
    if isinstance(value, datetime.datetime) and value.tzinfo:
        value = value.isoformat()
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
    return value


def write_report(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        headers = list(dict.fromkeys(k for row in rows for k in row)) or ["Information"]
        ws.append(headers)
        for row_number, row in enumerate(rows, 2):
            ws.append([cell_value(row.get(k)) for k in headers])
            status = row.get("Status")
            color = COLORS.get(status)
            if status == "MATCHED" and (
                row.get("Metadata Quality") != "GOOD"
                or row.get("Integrity") == "FAIL"
                or row.get("Archive Eligible") is False
            ):
                color = COLORS["WARNING"]
            for column_number in range(1, len(headers) + 1):
                c = ws.cell(row_number, column_number)
                if isinstance(c.value, str):
                    c.data_type = "s"  # Metadata must never become an Excel formula.
                if color:
                    c.fill = PatternFill("solid", fgColor=color)
                key = headers[c.column - 1]
                if (
                    c.value
                    and isinstance(c.value, str)
                    and (
                        key in ("Full Path", "Matched File Path")
                        or key.startswith("Alternative Candidate")
                        and not key.endswith("Score")
                    )
                ):
                    try:
                        c.hyperlink = Path(c.value).as_uri()
                        c.font = Font(color="0563C1", underline="single")
                    except ValueError:
                        pass
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D9E2F3")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, key in enumerate(headers, 1):
            from openpyxl.utils import get_column_letter

            ws.column_dimensions[get_column_letter(i)].width = (
                65
                if "Path" in key or "Candidate" in key and not key.endswith("Score")
                else 55
                if any(s in key for s in ("Reason", "Issues", "Notes", "Error"))
                else 28
            )
    path = Path(path)
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    wb.save(temporary)
    try:
        os.replace(temporary, path)
        destination = path
    except PermissionError:
        # Excel commonly locks an existing workbook on Windows. Preserve the
        # completed audit in a deterministic alternate file rather than
        # deleting/overwriting anything or losing the scan.
        destination = path.with_name(path.stem + "_new.xlsx")
        os.replace(temporary, destination)
        print(
            f"Could not replace locked workbook; wrote alternate report: {destination}. "
            "Close the old workbook and rerun to update the configured filename.",
            flush=True,
        )
    else:
        print(f"Workbook written: {destination}", flush=True)
    return destination
