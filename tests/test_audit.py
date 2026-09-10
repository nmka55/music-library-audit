from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from normalization import prepare_apple, prepare_offline
from matcher import match_all
from album_audit import audit_albums
from duplicate_detector import duplicates
from cache import Cache
from excel_report import write_report
from openpyxl import load_workbook

C = {"matched_threshold": 90, "review_threshold": 70, "duration_tolerance_seconds": 3}


def apple(**kwargs):
    r = {
        "Name": "Signal",
        "Artist": "Artist One & Artist Two",
        "Album": "Original",
        "Total Time": 140000,
        "Persistent ID": "TEST",
        "Track Number": 1,
    }
    r.update(kwargs)
    prepare_apple(r)
    return r


def offline(**kwargs):
    r = {
        "Title": "Signal (feat. Artist Two)",
        "Artist": "Artist One",
        "Album": "Greatest Hits",
        "Duration": 140.5,
        "Full Path": str(Path(__file__).resolve()),
        "Library Root": str(Path(__file__).resolve().parent),
        "Archive Eligible": True,
        "Integrity": "PASS (probe only)",
        "Artwork Present": True,
        "Track Number": 1,
        "Track Total": 2,
        "Metadata Issues": "",
    }
    r.update(kwargs)
    prepare_offline(r)
    return r


class SafetyTests(unittest.TestCase):
    def match(self, a=None, b=None, override=None):
        return match_all([a or apple()], [b or offline()], C, override or {})[0]

    def test_feature_credit_and_different_album(self):
        self.assertEqual(self.match()["Status"], "MATCHED")

    def test_comma_credits_and_with(self):
        a = apple(Name="Example", Artist="Artist One, Ártist Two & Artist Three")
        b = offline(Title="Example (with Artist Two & Artist Three)", Artist="Artist One")
        self.assertEqual(self.match(a=a, b=b)["Status"], "MATCHED")

    def test_unrelated_song_with_same_version_suffix(self):
        a = apple(Name="Mountain (Radio Edit)", Artist="Artist One")
        b = offline(Title="Ocean (Radio Edit)", Artist="Artist One", Duration=152)
        self.assertEqual(self.match(a=a, b=b)["Status"], "MISSING")

    def test_album_live_marker_and_numbered_parts(self):
        self.assertNotEqual(
            self.match(b=offline(Album="Live at Example Hall"))["Status"], "MATCHED"
        )
        a = apple(Name="An extremely long title for a recording part 1")
        b = offline(Title="An extremely long title for a recording part 2 (feat. Artist Two)")
        self.assertNotEqual(self.match(a=a, b=b)["Status"], "MATCHED")

    def test_versions_never_green(self):
        for version in (
            "Live",
            "Acoustic",
            "Demo",
            "Remix",
            "Remastered 2020",
            "Radio Edit",
            "Extended Mix",
            "Instrumental",
            "Karaoke",
            "Mono",
            "Stereo",
            "Edit",
            "Version",
        ):
            with self.subTest(version=version):
                r = self.match(b=offline(Title=f"Signal ({version}) (feat. Artist Two)"))
                self.assertNotEqual(r["Status"], "MATCHED")

    def test_missing_duration_or_artist_never_green(self):
        for changes in ({"Duration": None}, {"Artist": ""}, {"Title": ""}):
            self.assertNotEqual(self.match(b=offline(**changes))["Status"], "MATCHED")

    def test_protected_copy_cannot_satisfy_match(self):
        b = offline(**{"Archive Eligible": False, "DRM Status": "Protected / DRM"})
        self.assertEqual(self.match(b=b)["Status"], "REVIEW")
        self.assertEqual(self.match(b=b, override={"TEST": b["Full Path"]})["Status"], "REVIEW")

    def test_artwork_and_tags_separate_from_match(self):
        r = self.match(
            b=offline(
                **{"Artwork Present": False, "Metadata Issues": "Missing Artwork; Missing Album"}
            )
        )
        self.assertEqual(r["Status"], "MATCHED")
        self.assertIn("Missing Artwork", r["Metadata Quality"])

    def test_wrong_artist(self):
        self.assertNotEqual(self.match(b=offline(Artist="Unrelated Singer"))["Status"], "MATCHED")

    def test_duration_boundaries(self):
        self.assertEqual(self.match(b=offline(Duration=143))["Status"], "MATCHED")
        for duration in (143.1, 145, 155):
            self.assertNotEqual(self.match(b=offline(Duration=duration))["Status"], "MATCHED")

    def test_manual_missing_no_substitute(self):
        r = self.match(override={"TEST": str(Path(tempfile.gettempdir()) / "does-not-exist.flac")})
        self.assertEqual(r["Status"], "REVIEW")
        self.assertIn("does-not-exist", r["Matched File Path"])

    def test_manual_valid(self):
        r = self.match(override={"TEST": offline()["Full Path"]})
        self.assertEqual(r["Match Type"], "MANUAL MATCH")
        self.assertEqual(r["Overall Confidence"], 100)

    def test_top_three_and_safe_copy_preferred(self):
        bs = [
            offline(
                **{
                    "Full Path": str(Path(tempfile.gettempdir()) / f"{i}.flac"),
                    "Duration": 141 + i / 10,
                    "Archive Eligible": False,
                }
            )
            for i in range(4)
        ]
        r = match_all([apple()], bs, C, {})[0]
        self.assertEqual(r["Status"], "REVIEW")
        self.assertTrue(all(r[f"Alternative Candidate {i}"] for i in (1, 2, 3)))
        bs.append(offline(Duration=142))
        self.assertEqual(match_all([apple()], bs, C, {})[0]["Status"], "MATCHED")

    def test_album_holes_and_disc_holes(self):
        rows = [
            offline(**{"Disc Total": 2}),
            offline(**{"Track Number": 3, "Track Total": 3, "Disc Total": 2}),
        ]
        r = audit_albums(rows)[0]
        self.assertEqual(r["Status"], "INCONSISTENT TAGS")
        self.assertEqual(r["Missing Track Numbers"], [2])
        self.assertEqual(r["Missing Disc Numbers"], [2])

    def test_exact_duplicate(self):
        rows = [
            offline(**{"SHA-256": "abc"}),
            offline(
                **{"SHA-256": "abc", "Full Path": str(Path(tempfile.gettempdir()) / "copy.flac")}
            ),
        ]
        self.assertEqual(len(duplicates(rows)), 2)
        self.assertEqual(duplicates(rows)[0]["Category"], "EXACT BINARY DUPLICATE")

    def test_sidecars_do_not_count_as_duplicate_audio(self):
        rows = [
            offline(**{"SHA-256": "abc", "File Type": "AppleDouble metadata sidecar"}),
            offline(
                **{
                    "SHA-256": "abc",
                    "File Type": "AppleDouble metadata sidecar",
                    "Full Path": str(Path(tempfile.gettempdir()) / "._copy.m4a"),
                }
            ),
        ]
        self.assertEqual(duplicates(rows), [])

    def test_cache_invalidation(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "file"
            p.write_bytes(b"abc")
            cache = Cache(Path(td))
            cache.put(p, p.stat(), {"FFprobe OK": True, "Deep Integrity": "NOT TESTED"})
            self.assertIsNotNone(cache.get(p, p.stat(), False))
            self.assertIsNone(cache.get(p, p.stat(), True))
            p.write_bytes(b"changed")
            self.assertIsNone(cache.get(p, p.stat(), False))
            cache.close()

    def test_excel_formula_safety_and_warning_color(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "test.xlsx"
            write_report(
                p,
                {
                    "Match Results": [
                        {
                            "Title": "=1+1",
                            "Status": "MATCHED",
                            "Metadata Quality": "Missing Artwork",
                        }
                    ]
                },
            )
            wb = load_workbook(p)
            ws = wb.active
            self.assertEqual(ws["A2"].data_type, "s")
            self.assertEqual(ws["A2"].fill.fgColor.rgb, "00FFF2CC")
            self.assertEqual(ws.freeze_panes, "A2")
            wb.close()


if __name__ == "__main__":
    unittest.main()
