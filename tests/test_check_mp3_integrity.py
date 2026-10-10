from pathlib import Path
import plistlib
import shutil
import subprocess
import sys
import tempfile
import unittest

from mutagen.easyid3 import EasyID3
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import check_mp3_integrity as m
from config import DEFAULTS
from normalization import prepare_offline
from compare_libraries import as_file_rows
from track_matching import LooseIndex

HEADER = b"\xff\xfb\x90\x00"  # MPEG-1 Layer III, 128 kbps, 44.1 kHz, stereo: 417-byte frames
FRAME = HEADER + bytes(413)
ID3V2 = b"ID3\x03\x00\x00\x00\x00\x00\x00"
CONFIG = {**DEFAULTS, "_workspace": ".", "_config_path": "config.json"}


def info_frame(frames, size):
    frame = bytearray(FRAME)
    frame[36:40] = b"Info"
    frame[40:52] = (3).to_bytes(4, "big") + frames.to_bytes(4, "big") + size.to_bytes(4, "big")
    return bytes(frame)


def make_mp3(frames=200, with_info=True):
    body = FRAME * frames
    info = info_frame(frames, len(body) + 417) if with_info else b""
    return ID3V2 + info + body


class StructureScanTests(unittest.TestCase):
    def check(self, data, **kwargs):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "song.mp3"
            path.write_bytes(data)
            result = m.scan(path)
            return result, m.judge(result, **kwargs)

    def test_clean_file_with_id3_tags_and_trailing_padding_is_ok(self):
        data = make_mp3() + bytes(300) + b"TAG" + bytes(125)
        result, (verdict, issues) = self.check(data)
        self.assertEqual((verdict, issues), (m.OK, []))
        self.assertEqual((result["frames"], result["expected_frames"]), (200, 200))

    def test_bytes_cut_out_mid_frame_are_found_with_their_position(self):
        data = make_mp3()
        cut = len(ID3V2) + 417 * 101 + 100
        result, (verdict, issues) = self.check(data[:cut] + data[cut + 100 :])
        self.assertEqual(verdict, m.CORRUPT)
        self.assertEqual(len(result["gaps"]), 1)
        self.assertAlmostEqual(result["gaps"][0][2], 100 * 1152 / 44100, delta=0.1)

    def test_whole_frames_removed_are_found_through_the_length_header(self):
        data = make_mp3()
        start = len(ID3V2) + 417 * 60
        result, (verdict, issues) = self.check(data[:start] + data[start + 417 * 50 :])
        self.assertEqual((verdict, result["gaps"]), (m.CORRUPT, []))
        self.assertIn("50 frames missing", issues[0])

    def test_truncated_download_is_found(self):
        data = make_mp3()
        result, (verdict, issues) = self.check(data[: len(data) // 2 + 100])
        self.assertEqual(verdict, m.CORRUPT)
        self.assertTrue(result["partial_tail"])
        self.assertTrue(any("shorter than its header" in i for i in issues))

    def test_truncation_without_a_length_header_is_only_suspect(self):
        data = make_mp3(with_info=False)
        _, (verdict, issues) = self.check(data[:-100])
        self.assertEqual(verdict, m.SUSPECT)
        self.assertIn("middle of a frame", issues[0])

    def test_fake_headers_inside_audio_do_not_confuse_resync(self):
        noisy = HEADER + b"\xff\xfb\x90\x00" * 50 + bytes(413 - 200)
        body = (FRAME * 10) + noisy + (FRAME * 10)
        data = ID3V2 + info_frame(21, len(body) + 417) + body
        cut = len(ID3V2) + 417 * 5 + 20
        result, _ = self.check(data[:cut] + data[cut + 40 :])
        self.assertEqual(len(result["gaps"]), 1)
        self.assertLess(result["gaps"][0][1], 417 * 2)

    def test_unusable_files_are_reported_not_crashed_on(self):
        for data, text in (
            (b"", "empty"),
            (b"\x00\x00\x00\x18ftypM4A " + bytes(100), "MP4"),
            (b"hello" * 40, "no MP3"),
        ):
            _, (verdict, issues) = self.check(data)
            self.assertEqual(verdict, m.UNREADABLE)
            self.assertIn(text, issues[0])


class LibraryCrossCheckTests(unittest.TestCase):
    def reference(self, seconds):
        track = {
            "Name": "Song",
            "Artist": "Band",
            "Total Time": seconds * 1000,
            "Persistent ID": "1",
        }
        rows = as_file_rows([track])
        for row in rows:
            prepare_offline(row)
        return LooseIndex(rows)

    def test_file_much_shorter_than_the_library_track_is_suspect(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "song.mp3"
            path.write_bytes(make_mp3(with_info=False))  # about 5.2 s of audio
            tags = EasyID3()
            tags["title"], tags["artist"] = "Song", "Band"
            tags.save(path)
            seconds = m.scan(path)["seconds"]
            for official, expected in ((30, m.SUSPECT), (round(seconds), m.OK)):
                shortfall = m.library_shortfall(path, seconds, self.reference(official), CONFIG)
                verdict, issues = m.judge(m.scan(path), None, None, shortfall)
                self.assertEqual(verdict, expected, issues)


@unittest.skipUnless(shutil.which("ffmpeg"), "Requires FFmpeg")
class EndToEndTests(unittest.TestCase):
    def test_real_mp3_damage_report_workbook_and_read_only(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            music = base / "music"
            music.mkdir()
            clean = music / "clean.mp3"
            subprocess.run(
                [
                    "ffmpeg",
                    "-nostdin",
                    "-v",
                    "error",
                    "-f",
                    "lavfi",
                    "-i",
                    "anoisesrc=d=20:c=pink:r=44100",
                ]
                + ["-c:a", "libmp3lame", "-b:a", "128k", str(clean)],
                check=True,
            )
            data = clean.read_bytes()
            (music / "cut.mp3").write_bytes(data[:120000] + data[150000:])
            (music / "truncated.mp3").write_bytes(data[:200000])
            library = base / "library.xml"
            library.write_bytes(plistlib.dumps({"Tracks": {}}))
            before = {p.name: p.read_bytes() for p in music.iterdir()}
            code = m.main(
                [
                    "--workspace",
                    str(base / "work"),
                    "--music-root",
                    str(music),
                    "--decode",
                    "--all",
                ],
                tell=lambda _: None,
            )
            self.assertEqual(code, 0)
            self.assertEqual(before, {p.name: p.read_bytes() for p in music.iterdir()})
            sheet = load_workbook(base / "work" / "output" / "MP3_Integrity_Report.xlsx").active
            table = list(sheet.iter_rows(values_only=True))
            self.assertEqual((sheet.max_column, list(table[0])), (len(m.COLUMNS), list(m.COLUMNS)))
            verdicts = {r[1]: r[0] for r in table[1:]}
            self.assertEqual(
                verdicts, {"clean.mp3": m.OK, "cut.mp3": m.CORRUPT, "truncated.mp3": m.CORRUPT}
            )


if __name__ == "__main__":
    unittest.main()
